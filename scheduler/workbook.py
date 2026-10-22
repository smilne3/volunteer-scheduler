"""Opening signups.xlsx and finding sheets and columns, with clear errors."""

import zipfile
from pathlib import Path
from typing import Dict

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .errors import InputError
from .text import is_blank, norm


def open_workbook(path: Path):
    path = Path(path)
    if not path.exists():
        # iCloud keeps a hidden ".name.icloud" placeholder for files it hasn't downloaded.
        if path.with_name(f".{path.name}.icloud").exists():
            raise InputError(
                f"{path.name} is in iCloud but hasn't been downloaded to this Mac yet. "
                f"In Finder, right-click {path.name} and choose Download Now, then run this again."
            )
        raise InputError(
            f"Can't find {path}. Save the form export there as {path.name}, "
            f"or point to it with --file."
        )
    try:
        # Read-only: the tool never writes to the sign-up file (decision D2).
        return load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as e:
        raise InputError(
            f"{path.name} couldn't be opened as an Excel file ({e.__class__.__name__}). "
            f"Download it again from Google Sheets with File > Download > Microsoft Excel (.xlsx)."
        )


def get_sheet(wb, name: str):
    if name in wb.sheetnames:
        return wb[name]
    found = ", ".join(f'"{s}"' for s in wb.sheetnames)
    raise InputError(
        f'The workbook has no sheet named "{name}". Sheets found: {found}. '
        f'If it was renamed, rename it back to "{name}".'
    )


def find_columns(ws, sheet_name: str, required: Dict[str, str], optional: Dict[str, str] = None) -> Dict[str, int]:
    """Map each key to a 0-based column index by matching header text in row 1."""
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = {norm(h): i for i, h in enumerate(header_row) if not is_blank(h)}

    columns, missing = {}, []
    for key, title in required.items():
        if norm(title) in headers:
            columns[key] = headers[norm(title)]
        else:
            missing.append(title)
    for key, title in (optional or {}).items():
        if norm(title) in headers:
            columns[key] = headers[norm(title)]

    if missing:
        found = "\n".join(f'  - "{h}"' for h in header_row if not is_blank(h)) or "  (none)"
        wanted = "\n".join(f'  - "{m}"' for m in missing)
        raise InputError(
            f'The "{sheet_name}" sheet is missing these column headings:\n{wanted}\n'
            f"Headings found in row 1:\n{found}\n"
            f"If a question or column was renamed, change its heading back to match."
        )
    return columns
