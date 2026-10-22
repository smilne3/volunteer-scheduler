"""Phase 2: availability_review.csv, where the coordinator confirms free-text readings."""

import csv
import io
import shutil
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scheduler.errors import InputError
from scheduler.review_file import CONFIRMED, TO_REVIEW, apply_reviews
from scheduler.signups import read_signups

FIXTURE = Path(__file__).parent / "fixtures" / "signups_fake.xlsx"
TODAY = date(2026, 10, 22)


@pytest.fixture
def folder(tmp_path):
    shutil.copy(FIXTURE, tmp_path / "signups.xlsx")
    return tmp_path


def run(folder):
    data = read_signups(folder / "signups.xlsx")
    review = apply_reviews(data, today=TODAY)
    return data, review


def person(data, name):
    return next(v for v in data.volunteers if v.name == name)


def read_csv(path):
    return list(csv.DictReader(io.StringIO(path.read_bytes().decode("utf-8-sig"))))


def confirm(path, readings, encoding="utf-8-sig"):
    """Fill confirmed_reading for {name: reading}, the way the coordinator would."""
    rows = read_csv(path)
    for r in rows:
        if r["name"] in readings:
            r["confirmed_reading"] = readings[r["name"]]
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    path.write_bytes(buf.getvalue().encode(encoding))


def test_first_run_creates_file_with_every_free_text_answer(folder):
    data, review = run(folder)
    path = folder / "availability_review.csv"
    assert review.written and path.exists()
    rows = read_csv(path)
    assert len(rows) == 11
    assert all(r["status"] == TO_REVIEW and r["confirmed_reading"] == "" for r in rows)
    bailey = next(r for r in rows if r["name"] == "Bailey Sample")
    assert bailey["suggested_reading"] == "SAT-0900, SAT-1200 except 2026-11-14"
    assert bailey["answer"] == "Saturdays but not the 14th"


def test_unconfirmed_people_are_never_available(folder):
    data, review = run(folder)
    for item in review.items:
        assert item.volunteer.reading is None
        assert item.volunteer.status == "needs reading"
    assert sum(v.status == "ready" for v in data.volunteers) == 19


def test_confirmed_reading_is_used(folder):
    run(folder)
    confirm(folder / "availability_review.csv", {"Bailey Sample": "SAT-0900, SAT-1200 except 2026-11-14"})
    data, review = run(folder)
    bailey = person(data, "Bailey Sample")
    assert bailey.status == "ready" and bailey.reading_source == "confirmed"
    assert not bailey.reading.allows("SAT-0900", date(2026, 11, 14))
    assert bailey.reading.allows("SAT-0900", date(2026, 11, 7))
    assert review.count(CONFIRMED) == 1 and review.count(TO_REVIEW) == 10
    # The file keeps the confirmation and shows it as confirmed.
    row = next(r for r in read_csv(folder / "availability_review.csv") if r["name"] == "Bailey Sample")
    assert row["confirmed_reading"] == "SAT-0900, SAT-1200 except 2026-11-14"
    assert row["status"] == CONFIRMED


def test_a_mistake_in_a_confirmed_reading_keeps_them_off_the_schedule(folder):
    run(folder)
    confirm(folder / "availability_review.csv", {"Sam Prototype": "THU-1500 except 2026-11-27"})
    data, review = run(folder)
    assert person(data, "Sam Prototype").reading is None
    row = next(r for r in read_csv(folder / "availability_review.csv") if r["name"] == "Sam Prototype")
    assert row["status"].startswith("error: 2026-11-27 is a Friday")
    assert row["confirmed_reading"] == "THU-1500 except 2026-11-27"   # their typing is kept


def test_confirmed_none_means_no_shifts(folder):
    run(folder)
    confirm(folder / "availability_review.csv", {"Dee Consectetur": "none"})
    data, _ = run(folder)
    dee = person(data, "Dee Consectetur")
    assert dee.status == "no shifts" and dee.reading.shift_ids == frozenset()


def test_excel_style_save_still_matches(folder):
    """Excel may save as Windows-1252 and rewrite the timestamps."""
    path = folder / "availability_review.csv"
    run(folder)
    rows = read_csv(path)
    for r in rows:
        r["response_time"] = "10/17/26 15:36"
        if r["name"] == "Val Variable":
            r["confirmed_reading"] = "SAT-0900; WED-1700 only 2026-10-28, 2026-11-11"
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    path.write_bytes(buf.getvalue().encode("cp1252"))   # Val's answer has an en dash
    data, _ = run(folder)
    val = person(data, "Val Variable")
    assert val.reading_source == "confirmed"
    assert val.reading.allows("WED-1700", date(2026, 11, 11))
    assert not val.reading.allows("WED-1700", date(2026, 11, 4))


def add_response(folder, name, email, available):
    wb = load_workbook(folder / "signups.xlsx")
    wb["Form Responses 1"].append([datetime(2026, 10, 22, 9), name, email, "(312) 555-0102", available, "", 1, "No"])
    wb.save(folder / "signups.xlsx")


def test_new_response_with_a_different_answer_needs_confirming_again(folder):
    run(folder)
    confirm(folder / "availability_review.csv", {"Bailey Sample": "SAT-0900, SAT-1200 except 2026-11-14"})
    add_response(folder, "Bailey Sample", "bailey.sample@example.com", "Saturdays, but not Nov 21")
    data, review = run(folder)
    bailey = person(data, "Bailey Sample")
    assert bailey.reading is None and bailey.status == "needs reading"
    assert [r["confirmed_reading"] for r in review.no_longer_apply] == ["SAT-0900, SAT-1200 except 2026-11-14"]


def test_new_response_with_the_same_answer_keeps_the_confirmation(folder):
    run(folder)
    confirm(folder / "availability_review.csv", {"Bailey Sample": "SAT-0900, SAT-1200 except 2026-11-14"})
    add_response(folder, "Bailey Sample", "Bailey.Sample@example.com", "Saturdays but not the 14th")
    data, review = run(folder)
    assert person(data, "Bailey Sample").reading_source == "confirmed"
    assert not review.no_longer_apply


def test_rerun_without_changes_does_not_rewrite(folder):
    run(folder)
    _, review = run(folder)
    assert not review.written


def test_file_open_in_excel_is_not_overwritten(folder):
    run(folder)
    path = folder / "availability_review.csv"
    confirm(path, {"Kai Madeup": "SAT-0900"})
    before = path.read_bytes()
    (folder / "~$availability_review.csv").write_bytes(b"")
    data, review = run(folder)
    assert path.read_bytes() == before
    assert "looks open in Excel" in review.warning
    assert person(data, "Kai Madeup").reading_source == "confirmed"   # still read


def test_missing_column_is_a_clear_error(folder):
    (folder / "availability_review.csv").write_text("name,email,answer\nA,a@example.com,x\n")
    with pytest.raises(InputError) as e:
        run(folder)
    assert "missing the column(s): confirmed_reading" in str(e.value)


def test_no_free_text_means_no_file(tmp_path):
    wb = load_workbook(FIXTURE)
    ws = wb["Form Responses 1"]
    for row in range(ws.max_row, 1, -1):
        if ws.cell(row, 1).value and ws.cell(row, 2).value != "Alex Example":
            ws.delete_rows(row)
    wb.save(tmp_path / "signups.xlsx")
    data = read_signups(tmp_path / "signups.xlsx")
    review = apply_reviews(data, today=TODAY)
    assert not review.items and not (tmp_path / "availability_review.csv").exists()


# --- preferences -------------------------------------------------------------

def test_free_text_preference_is_read_within_confirmed_availability(folder):
    run(folder)
    confirm(folder / "availability_review.csv", {"Bailey Sample": "SAT-0900, SAT-1200 except 2026-11-14"})
    data, review = run(folder)
    bailey = person(data, "Bailey Sample")            # "Mornings if possible"
    assert bailey.preferred_ids == {"SAT-0900"} and bailey.preferred_source == "suggested"


def test_preference_outside_availability_is_dropped(folder):
    run(folder)
    confirm(folder / "availability_review.csv", {"Dee Consectetur": "SAT-0900"})
    data, _ = run(folder)
    assert person(data, "Dee Consectetur").preferred_ids == frozenset()   # "afternoon"


def test_unconfirmed_people_have_no_preferences(folder):
    data, _ = run(folder)
    assert person(data, "Harper Demo").preferred_ids == frozenset()        # "tutoring!!"


def test_checkbox_preferences_still_come_from_the_form(folder):
    data, _ = run(folder)
    alex = person(data, "Alex Example")
    assert alex.preferred_ids == {"SAT-0900"} and alex.preferred_source == "form"
