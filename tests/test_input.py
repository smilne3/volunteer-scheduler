"""Phase 1: reading signups.xlsx.

tests/fixtures/signups_fake.xlsx is made-up data (fake names, 555 numbers,
example.com emails). Never put real sign-ups in tests/fixtures: this repo is public.
"""

import hashlib
from datetime import datetime, time
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from scheduler.errors import InputError
from scheduler.shifts import parse_time_range
from scheduler.signups import CLEAN, NEEDS_READING, read_signups

FIXTURE = Path(__file__).parent / "fixtures" / "signups_fake.xlsx"
SAT_AM = "Saturday Food Pantry 9am–12pm"
WED = "Wednesday Phone Bank 5pm–7pm"


def edited_copy(tmp_path, edit):
    """Copy the fake workbook, apply edit(wb), save, return the new path."""
    wb = load_workbook(FIXTURE)
    edit(wb)
    out = tmp_path / "signups.xlsx"
    wb.save(out)
    return out


def add_response(wb, ts, name, email, available, preferred="", limit=1, before="No"):
    wb["Form Responses 1"].append([ts, name, email, "(312) 555-0199", available, preferred, limit, before])


def error_text(path) -> str:
    with pytest.raises(InputError) as e:
        read_signups(path)
    return str(e.value)


def by_name(data, name):
    return next(v for v in data.volunteers if v.name == name)


# --- the fake data -------------------------------------------------------------

def test_fake_data_reads_as_expected():
    data = read_signups(FIXTURE)
    assert [s.id for s in data.shifts] == ["TUE-1500", "WED-1700", "THU-1500", "SAT-0900", "SAT-1200"]
    assert len(data.volunteers) == 30
    statuses = [v.status for v in data.volunteers]
    assert statuses.count("ready") == 19
    assert statuses.count("needs reading") == 11
    assert not data.replaced and not data.warnings


def test_clean_answer_becomes_shift_ids():
    alex = by_name(read_signups(FIXTURE), "Alex Example")
    assert alex.available.status == CLEAN
    assert alex.available.shift_ids == {"SAT-0900", "SAT-1200"}
    assert alex.preferred.shift_ids == {"SAT-0900"}
    assert alex.max_per_week == 2 and alex.experienced
    assert alex.email == "alex.example@example.com"
    assert isinstance(alex.timestamp, datetime)


def test_partly_matching_free_text_still_needs_reading():
    # "Saturday Food Pantry 9am–12pm, and phone bank every other week"
    val = by_name(read_signups(FIXTURE), "Val Variable")
    assert val.available.status == NEEDS_READING
    assert val.available.shift_ids == frozenset()


def test_shifts_sheet_details():
    shifts = {s.id: s for s in read_signups(FIXTURE).shifts}
    sat = shifts["SAT-0900"]
    assert (sat.day, sat.start, sat.end, sat.needs, sat.needs_experienced) == (
        "Saturday", time(9), time(12), 3, True)
    assert shifts["TUE-1500"].needs_experienced is False


def test_reading_never_changes_the_file():
    before = hashlib.sha256(FIXTURE.read_bytes()).hexdigest()
    read_signups(FIXTURE)
    assert hashlib.sha256(FIXTURE.read_bytes()).hexdigest() == before


# --- answer matching is forgiving about dashes, case and spacing ----------------

def test_answer_matching_ignores_dash_case_and_spacing(tmp_path):
    path = edited_copy(tmp_path, lambda wb: add_response(
        wb, datetime(2026, 10, 21, 9), "Lou Lenient", "lou@example.com",
        "saturday food pantry 9am - 12pm,   WEDNESDAY PHONE BANK 5PM-7PM"))
    lou = by_name(read_signups(path), "Lou Lenient")
    assert lou.available.status == CLEAN
    assert lou.available.shift_ids == {"SAT-0900", "WED-1700"}


def test_preference_outside_availability_is_noted(tmp_path):
    path = edited_copy(tmp_path, lambda wb: add_response(
        wb, datetime(2026, 10, 21, 9), "Pip Picky", "pip@example.com", SAT_AM, WED))
    pip = by_name(read_signups(path), "Pip Picky")
    assert pip.status == "ready"
    assert any("WED-1700" in n and "ignored" in n for n in pip.notes)


# --- file-level errors stop the run ------------------------------------------

def test_missing_file(tmp_path):
    assert "Can't find" in error_text(tmp_path / "signups.xlsx")


def test_file_in_icloud_but_not_downloaded(tmp_path):
    (tmp_path / ".signups.xlsx.icloud").write_bytes(b"")
    msg = error_text(tmp_path / "signups.xlsx")
    assert "hasn't been downloaded" in msg and "Download Now" in msg


def test_not_an_excel_file(tmp_path):
    path = tmp_path / "signups.xlsx"
    path.write_text("Timestamp,Full name\n")
    assert "couldn't be opened as an Excel file" in error_text(path)


def test_renamed_responses_sheet(tmp_path):
    def rename(wb):
        wb["Form Responses 1"].title = "Responses"
    msg = error_text(edited_copy(tmp_path, rename))
    assert 'no sheet named "Form Responses 1"' in msg and '"Responses"' in msg


def test_renamed_shifts_sheet(tmp_path):
    def rename(wb):
        wb["Shifts"].title = "Shift list"
    assert 'no sheet named "Shifts"' in error_text(edited_copy(tmp_path, rename))


def test_renamed_column(tmp_path):
    def rename(wb):
        wb["Form Responses 1"]["G1"] = "Max shifts"
    msg = error_text(edited_copy(tmp_path, rename))
    assert "What is the most shifts you'd like per week?" in msg
    assert '"Max shifts"' in msg  # lists what it did find


def test_column_heading_case_and_curly_apostrophe_still_match(tmp_path):
    def restyle(wb):
        wb["Form Responses 1"]["G1"] = "WHAT IS THE MOST SHIFTS YOU’D LIKE PER WEEK?"
    assert len(read_signups(edited_copy(tmp_path, restyle)).volunteers) == 30


def test_shift_name_missing_from_shifts_sheet_stops_the_run(tmp_path):
    def add(wb):
        for i, name in enumerate(["Fay Friday", "Gus Friday"]):
            add_response(wb, datetime(2026, 10, 21, 9 + i), name, f"{name.split()[0].lower()}@example.com",
                         f"{SAT_AM}, Friday Food Pantry 9am–12pm")
    msg = error_text(edited_copy(tmp_path, add))
    assert "Friday Food Pantry 9am–12pm" in msg
    assert "rows 32, 33" in msg
    assert "aren't in the Shifts sheet" in msg


def test_one_person_typing_a_shift_like_answer_is_only_a_warning(tmp_path):
    path = edited_copy(tmp_path, lambda wb: add_response(
        wb, datetime(2026, 10, 21, 9), "Ty Typer", "ty@example.com", "Tuesday 3-5pm"))
    data = read_signups(path)
    assert by_name(data, "Ty Typer").status == "needs reading"
    assert any("Tuesday 3-5pm" in w for w in data.warnings)


# --- row-level problems are listed and only block that person ------------------

@pytest.mark.parametrize("limit, expected", [
    (None, "No answer for the most shifts per week"),
    ("two", 'is "two", which isn\'t a whole number'),
    (1.5, 'is "1.5", which isn\'t a whole number'),
    (-1, 'is "-1", which isn\'t a whole number'),
])
def test_bad_weekly_limit_is_a_problem_for_that_person(tmp_path, limit, expected):
    path = edited_copy(tmp_path, lambda wb: add_response(
        wb, datetime(2026, 10, 21, 9), "Lim Itless", "lim@example.com", SAT_AM, limit=limit))
    data = read_signups(path)
    lim = by_name(data, "Lim Itless")
    assert lim.status == "problem"
    assert any(expected in p for p in lim.problems)
    assert len(data.volunteers) == 31  # everyone else still read


@pytest.mark.parametrize("limit, expected", [("3", 3), (2.0, 2), (0, 0)])
def test_weekly_limit_forms_that_are_fine(tmp_path, limit, expected):
    path = edited_copy(tmp_path, lambda wb: add_response(
        wb, datetime(2026, 10, 21, 9), "Lim Itless", "lim@example.com", SAT_AM, limit=limit))
    lim = by_name(read_signups(path), "Lim Itless")
    assert lim.max_per_week == expected and not lim.problems


def test_missing_email_is_a_problem(tmp_path):
    path = edited_copy(tmp_path, lambda wb: add_response(
        wb, datetime(2026, 10, 21, 9), "No Mail", "", SAT_AM))
    assert by_name(read_signups(path), "No Mail").status == "problem"


def test_newest_response_wins_by_timestamp_not_row(tmp_path):
    def add(wb):
        # Newer response sits in the earlier row (as if the sheet was re-sorted).
        add_response(wb, datetime(2026, 10, 21, 12), "Dee Dupe", "Dee@Example.com", WED, limit=2)
        add_response(wb, datetime(2026, 10, 20, 12), "Dee Dupe", "dee@example.com", SAT_AM, limit=1)
    data = read_signups(edited_copy(tmp_path, add))
    dees = [v for v in data.volunteers if v.email == "dee@example.com"]
    assert len(dees) == 1
    assert dees[0].row == 32 and dees[0].available.shift_ids == {"WED-1700"}
    assert [(r.old_row, r.new_row) for r in data.replaced] == [(33, 32)]


def test_blank_rows_are_skipped(tmp_path):
    def add(wb):
        wb["Form Responses 1"].append([None] * 8)
        add_response(wb, datetime(2026, 10, 21, 9), "After Blank", "ab@example.com", SAT_AM)
    data = read_signups(edited_copy(tmp_path, add))
    assert len(data.volunteers) == 31


# --- Shifts sheet problems ---------------------------------------------------

def shifts_sheet_copy(tmp_path, rows):
    """Fake workbook with the Shifts sheet replaced by these rows."""
    def replace(wb):
        del wb["Shifts"]
        ws = wb.create_sheet("Shifts")
        ws.append(["Shift", "Day", "Time", "People needed", "Needs at least one experienced volunteer?"])
        for r in rows:
            ws.append(r)
    return edited_copy(tmp_path, replace)


GOOD_ROWS = [
    ["Saturday Food Pantry 9am–12pm", "Saturday", "9:00 AM – 12:00 PM", 3, "Yes"],
    ["Saturday Food Pantry 12pm–3pm", "Saturday", "12:00 PM – 3:00 PM", 3, "Yes"],
    ["Tuesday Tutoring 3pm–5pm", "Tuesday", "3:00 PM – 5:00 PM", 2, "No"],
    ["Wednesday Phone Bank 5pm–7pm", "Wednesday", "5:00 PM – 7:00 PM", 3, "Yes"],
    ["Thursday Tutoring 3pm–5pm", "Thursday", "3:00 PM – 5:00 PM", 2, "No"],
]


def test_shifts_sheet_lists_every_bad_row_at_once(tmp_path):
    rows = [list(r) for r in GOOD_ROWS]
    rows[0][2] = "morning"         # row 2: unreadable time
    rows[1][3] = "three"           # row 3: needs not a number
    rows[2][4] = "Maybe"           # row 4: not yes/no
    rows[3][1] = "Wendsday"        # row 5: not a day
    rows[4][2] = None              # row 6: incomplete
    msg = error_text(shifts_sheet_copy(tmp_path, rows))
    assert 'Row 2 (Saturday Food Pantry 9am–12pm): Time "morning" isn\'t readable' in msg
    assert 'Row 3' in msg and 'People needed "three"' in msg
    assert 'Row 4' in msg and '"Maybe"; use Yes or No' in msg
    assert 'Row 5' in msg and 'Day "Wendsday"' in msg
    assert "Row 6: missing Time" in msg


def test_two_shifts_same_day_and_start(tmp_path):
    rows = GOOD_ROWS + [["Saturday Garden 9am–11am", "Sat", "9am - 11am", 2, "No"]]
    msg = error_text(shifts_sheet_copy(tmp_path, rows))
    assert "Rows 2 and 7 are both Saturday shifts starting at 9:00 AM" in msg


def test_shift_name_with_comma(tmp_path):
    rows = GOOD_ROWS + [["Friday Pantry, Setup", "Friday", "9am-11am", 2, "No"]]
    assert "contains a comma" in error_text(shifts_sheet_copy(tmp_path, rows))


def test_note_rows_in_shifts_sheet_are_ignored(tmp_path):
    rows = GOOD_ROWS + [[None] * 5, ["Note: these values are placeholders.", None, None, None, None]]
    assert len(read_signups(shifts_sheet_copy(tmp_path, rows)).shifts) == 5


def test_empty_shifts_sheet(tmp_path):
    assert "No shifts found" in error_text(shifts_sheet_copy(tmp_path, []))


@pytest.mark.parametrize("text, start, end", [
    ("9:00 AM – 12:00 PM", time(9), time(12)),
    ("12:00 PM - 3:00 PM", time(12), time(15)),
    ("3-5pm", time(15), time(17)),
    ("9-12pm", time(9), time(12)),
    ("5 to 7 pm", time(17), time(19)),
    ("10:30am–1pm", time(10, 30), time(13)),
])
def test_time_ranges(text, start, end):
    assert parse_time_range(text) == (start, end)


@pytest.mark.parametrize("text", ["morning", "5pm-3pm", "13:00 PM - 2 PM", ""])
def test_bad_time_ranges(text):
    assert parse_time_range(text) is None


def test_workbook_built_from_scratch_without_optional_phone_column(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Form Responses 1"
    ws.append(["Timestamp", "Full name", "Email address", "Which shifts are you available for?",
               "Which shifts would you prefer?", "What is the most shifts you'd like per week?",
               "Have you volunteered with us before?"])
    ws.append([datetime(2026, 10, 1, 9), "Solo Person", "solo@example.com", "Tue tutoring 3-5", "", 1, "Yes"])
    shifts = wb.create_sheet("Shifts")
    shifts.append(["Shift", "Day", "Time", "People needed", "Needs at least one experienced volunteer?"])
    shifts.append(["Tue tutoring 3-5", "Tue", "3-5pm", 2, "no"])
    path = tmp_path / "signups.xlsx"
    wb.save(path)
    data = read_signups(path)
    solo = data.volunteers[0]
    assert solo.status == "ready" and solo.phone == ""
    assert solo.available.shift_ids == {"TUE-1500"}
